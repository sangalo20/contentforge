# -*- coding: utf-8 -*-
# Template generator for Joe Downs short-form talking-topics Word doc.
# Copy this, swap the title-block strings + the `topics` list, set `out`, rerun with
#   C:/Python313/python.exe gen_shortform_doc.py
# Palette + layout are the approved Ep11 look (Storage Moguls brand).
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x0B, 0x13, 0x22)
BLUE = RGBColor(0x2E, 0x7D, 0xF6)
GREY = RGBColor(0x55, 0x5B, 0x66)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def heading(text, size=15, color=NAVY, space_before=16, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = color
    return p

def label(text, color=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = color
    return p

def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    return p

# ---- Title block (SWAP per episode) ----
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = t.add_run("Storage Moguls  |  EpNN Short-Form Recording Topics")
r.bold = True; r.font.size = Pt(20); r.font.color.rgb = NAVY

sub = doc.add_paragraph()
r = sub.add_run("Guest: <name, role>  \u00b7  Episode: \u201c<title>\u201d")
r.italic = True; r.font.size = Pt(10.5); r.font.color.rgb = GREY

note = doc.add_paragraph(); note.paragraph_format.space_after = Pt(6)
r = note.add_run("For Instagram, TikTok and Facebook (same content to all three). Joe records each one freely \u2014 the context line opens it so a cold viewer knows what\u2019s being talked about, then Joe builds to the point, then the CTA earns the click.")
r.font.size = Pt(10.5); r.font.color.rgb = GREY

d = doc.add_paragraph(); d.add_run("_" * 60).font.color.rgb = RGBColor(0xCC, 0xD2, 0xDA)

# ---- Topics (SWAP per episode) ----
# Each: title = PLAIN name (no clever title). context = cold-stranger setup FIRST.
# talk = build to the real insight from the Understanding. cta = reason -> YouTube -> storagemoguls.ai.
topics = [
    {
        "n": "1",
        "title": "<plain name of what this topic IS>",
        "context": "<cold-stranger setup: open on something a total stranger already understands, 2-3 sentences>",
        "talk": "<the build + the real insight pulled from the episode Understanding>",
        "cta": "<reason/payoff tied to THIS topic + full episode on YouTube + warm storagemoguls.ai invite>",
    },
]

for t in topics:
    heading(f"Topic {t['n']}: {t['title']}")
    label("CONTEXT (open with this)"); body(t["context"])
    label("WHAT TO TALK ABOUT"); body(t["talk"])
    label("CTA (finish strong)"); body(t["cta"])
    dv = doc.add_paragraph(); dv.paragraph_format.space_before = Pt(6)
    dv.add_run("_" * 60).font.color.rgb = RGBColor(0xE2, 0xE6, 0xEC)

out = r"C:\Users\somea\SM_EpNN\SM_EpNN_ShortForm_Topics.docx"
doc.save(out)
print("SAVED", out)
