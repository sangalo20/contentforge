"""Dump every slide of a .pptx to Markdown (text, tables, pictures, notes) for review.

    ../odin-agent/venv/bin/python deck_to_md.py human-is-the-reviewer.pptx > deck-review.md
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def shape_md(sh):
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "*(picture)*"
    if sh.has_table:
        t = sh.table
        rows = [[c.text.replace("\n", " ") for c in r.cells] for r in t.rows]
        out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
        out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        return "\n".join(out)
    if not sh.has_text_frame:
        return ""
    paras = [p.text for p in sh.text_frame.paragraphs]
    if not any(p.strip() for p in paras):
        return ""
    is_code = sh.text_frame.paragraphs[0].runs and sh.text_frame.paragraphs[0].runs[0].font.name == "Roboto Mono"
    if is_code:
        return "```\n" + "\n".join(paras) + "\n```"
    bulleted = sh._element.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}buChar")
    if bulleted and len(paras) > 1:
        return "\n".join(f"- {p}" for p in paras if p.strip())
    return "\n".join(f"**{p}**" if i == 0 and len(p) < 120 else p for i, p in enumerate(paras) if p.strip())


def main(path):
    prs = Presentation(path)
    print(f"# Deck review: {Path(path).name}\n")
    for i, s in enumerate(prs.slides, 1):
        print(f"\n---\n\n## Slide {i}\n")
        for sh in s.shapes:
            md = shape_md(sh)
            if md:
                print(md + "\n")
        if s.has_notes_slide:
            n = s.notes_slide.notes_text_frame.text.strip()
            if n:
                print("> **Speaker notes**")
                for line in n.split("\n"):
                    print(f"> {line}")
                print()


if __name__ == "__main__":
    main(sys.argv[1])
