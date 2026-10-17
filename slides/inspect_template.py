"""Dump template structure: per slide -> layout name, shapes (type, name, placeholder idx, text)."""
import sys
from pptx import Presentation
from pptx.util import Emu

prs = Presentation(sys.argv[1])
print("slide size", Emu(prs.slide_width).inches, Emu(prs.slide_height).inches)
print("layouts:", [l.name for l in prs.slide_layouts])
for i, s in enumerate(prs.slides, 1):
    print(f"\n--- slide {i}  layout={s.slide_layout.name!r}")
    for sh in s.shapes:
        ph = f"ph[{sh.placeholder_format.idx}:{sh.placeholder_format.type}]" if sh.is_placeholder else ""
        txt = sh.text_frame.text.replace("\n", " / ")[:70] if sh.has_text_frame else ""
        pos = f"({Emu(sh.left).inches:.1f},{Emu(sh.top).inches:.1f} {Emu(sh.width).inches:.1f}x{Emu(sh.height).inches:.1f})"
        print(f"  {sh.shape_type!s:22} {sh.name[:28]:28} {ph:14} {pos:22} {txt}")
